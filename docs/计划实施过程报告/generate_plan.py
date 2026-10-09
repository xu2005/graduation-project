# -*- coding: utf-8 -*-
"""《AI 辅助工程学习方案》docx 生成脚本
R4 封面(Top Color Block)+ GO-1 配色 + Profile A 正式排版
三节结构:封面(无页码)/ 摘要+目录(罗马数字)/ 正文(阿拉伯数字,从 1 起)
"""
import os
from docx import Document
from docx.shared import Pt, Twips, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"D:\ai data\z code\.zcode\workspace\default\AI辅助工程学习方案.docx"

# ── GO-1 Graphite Orange 配色 ──
ACCENT = "D4875A"        # 强调色(分隔线、表头底色)
DARK = "1A2330"          # 封面深色块
HEAD = "1A2330"          # 正文标题色
TBL_HEADER_BG = "D4875A"
TBL_HEADER_TX = "FFFFFF"
TBL_INNER = "DDD0C8"     # 表内横线
GRAY = "808080"
META = "4A5460"          # 封面白区信息文字

CONTENT_W = 8788         # 正文可用宽度(11906-1701-1417)


# ────────────────────────── 基础工具 ──────────────────────────
def _is_cjk(ch):
    return bool(ch) and ("\u4e00" <= ch <= "\u9fff")


_CJK_PUNCT_MAP = {":": "\uff1a", ";": "\uff1b", "?": "\uff1f", "!": "\uff01",
                  ",": "\uff0c", "(": "\uff08", ")": "\uff09"}


def cjk_normalize(text):
    """中文语境下的半角标点转全角:仅当标点紧邻汉字时转换,不影响纯 ASCII 内容"""
    if not text:
        return text
    chars = list(text)
    out = []
    for idx, ch in enumerate(chars):
        if ch in _CJK_PUNCT_MAP:
            prev = chars[idx - 1] if idx > 0 else ""
            nxt = chars[idx + 1] if idx < len(chars) - 1 else ""
            if _is_cjk(prev) or _is_cjk(nxt):
                out.append(_CJK_PUNCT_MAP[ch])
                continue
        out.append(ch)
    return "".join(out)


def set_run(run, size=12, bold=False, italic=False, color="000000",
            ascii_font="Times New Roman", east_font="SimSun"):
    if run.text:
        run.text = cjk_normalize(run.text)
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east_font)
    return run


def para(parent, text="", size=12, bold=False, italic=False, color="000000",
         ascii_font="Times New Roman", east_font="SimSun", align=None,
         first_indent=0, before=0, after=0, line=1.3, keep_next=False,
         line_rule=None):
    p = parent.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    pf.space_before = Twips(before)
    pf.space_after = Twips(after)
    if line_rule is not None:
        pf.line_spacing = Twips(line)
        pf.line_spacing_rule = line_rule
    else:
        pf.line_spacing = line
    if first_indent:
        pf.first_line_indent = Twips(first_indent)
    pf.keep_with_next = keep_next
    if text:
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, italic=italic, color=color,
                ascii_font=ascii_font, east_font=east_font)
    return p


def body(doc, text):
    return para(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                first_indent=480, after=80)


def heading(doc, text, level):
    p = doc.add_paragraph(style="Heading %d" % level)
    p.add_run(cjk_normalize(text))
    return p


def config_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string("000000")
    rPr = normal.element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")
    rf.set(qn("w:eastAsia"), "SimSun")
    npf = normal.paragraph_format
    npf.line_spacing = 1.3
    npf.space_before = Twips(0)
    npf.space_after = Twips(0)

    specs = [
        ("Heading 1", 16, 360, 200, True),
        ("Heading 2", 15, 280, 140, False),
        ("Heading 3", 14, 220, 110, False),
    ]
    for name, sz, before, after, centered in specs:
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(sz)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor.from_string(HEAD)
        rPr = st.element.get_or_add_rPr()
        rf = rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rPr.insert(0, rf)
        rf.set(qn("w:ascii"), "Times New Roman")
        rf.set(qn("w:hAnsi"), "Times New Roman")
        rf.set(qn("w:eastAsia"), "SimHei")
        hpf = st.paragraph_format
        hpf.space_before = Twips(before)
        hpf.space_after = Twips(after)
        hpf.line_spacing = 1.3
        hpf.keep_with_next = True
        if centered:
            hpf.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ────────────────────────── 表格工具 ──────────────────────────
def cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_width_pct(cell, pct):
    tcPr = cell._tc.get_or_add_tcPr()
    w = OxmlElement("w:tcW")
    w.set(qn("w:w"), str(int(pct * 50)))
    w.set(qn("w:type"), "pct")
    tcPr.append(w)


def row_cant_split(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))
    if header:
        trPr.append(OxmlElement("w:tblHeader"))


def table_width_pct(table, pct=100):
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(pct * 50))
    tblW.set(qn("w:type"), "pct")


def table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag, style, sz, color in [
        ("top", "single", 6, ACCENT),
        ("left", "none", 0, "auto"),
        ("bottom", "single", 6, ACCENT),
        ("right", "none", 0, "auto"),
        ("insideH", "single", 2, TBL_INNER),
        ("insideV", "none", 0, "auto"),
    ]:
        e = OxmlElement("w:" + tag)
        e.set(qn("w:val"), style)
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.append(borders)


def table_no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        e = OxmlElement("w:" + tag)
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        e.set(qn("w:color"), "auto")
        borders.append(e)
    tblPr.append(borders)


def table_cell_margins(table, top=60, bottom=60, left=120, right=120):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in [("top", top), ("left", left), ("bottom", bottom), ("right", right)]:
        e = OxmlElement("w:" + tag)
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def set_grid(table, pcts):
    grid = table._tbl.find(qn("w:tblGrid"))
    cols = grid.findall(qn("w:gridCol"))
    for c, p in zip(cols, pcts):
        c.set(qn("w:w"), str(int(CONTENT_W * p / 100)))


def row_keep_next(row):
    """对一行内所有段落设置 keep_with_next,使该行与下一行保持同页"""
    for c in row.cells:
        for pp in c.paragraphs:
            pp.paragraph_format.keep_with_next = True


def add_data_table(doc, caption, col_pcts, header, rows, center_cols=(), size=10.5):
    para(doc, caption, size=10.5, bold=True, east_font="SimHei",
         align=WD_ALIGN_PARAGRAPH.CENTER, before=140, after=80, keep_next=True)
    t = doc.add_table(rows=len(rows) + 1, cols=len(col_pcts))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    table_width_pct(t)
    table_borders(t)
    table_cell_margins(t)
    set_grid(t, col_pcts)

    hr = t.rows[0]
    row_cant_split(hr, header=True)
    row_keep_next(hr)  # 表头至少与第一行数据同页,防止页底孤立表头
    for j, txt in enumerate(header):
        c = hr.cells[j]
        cell_width_pct(c, col_pcts[j])
        cell_shading(c, TBL_HEADER_BG)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(txt)
        set_run(r, size=size, bold=True, color=TBL_HEADER_TX, east_font="SimHei")

    n = len(rows)
    keep_all = n <= 8  # 小表整表同页,大表依赖表头重复
    for i, rowdata in enumerate(rows, start=1):
        tr = t.rows[i]
        row_cant_split(tr)
        if (keep_all and i < n) or (not keep_all and i == 1):
            row_keep_next(tr)
        for j, txt in enumerate(rowdata):
            c = tr.cells[j]
            cell_width_pct(c, col_pcts[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            first = True
            for ln in str(txt).split("\n"):
                p = c.paragraphs[0] if first else c.add_paragraph()
                first = False
                p.paragraph_format.line_spacing = 1.3
                if j in center_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(ln)
                set_run(r, size=size)
    return t


def phase_table(doc, caption, items):
    """阶段实施要点表:2 列(项目 / 内容)"""
    add_data_table(doc, caption, [18, 82],
                   ["项目", "内容"], [[k, v] for k, v in items])


# ────────────────────────── 页码 / 页眉 ──────────────────────────
def set_pgnum(section, fmt, start):
    sectPr = section._sectPr
    el = sectPr.find(qn("w:pgNumType"))
    if el is None:
        el = OxmlElement("w:pgNumType")
        cols = sectPr.find(qn("w:cols"))
        if cols is not None:
            cols.addprevious(el)
        else:
            sectPr.append(el)
    el.set(qn("w:fmt"), fmt)
    el.set(qn("w:start"), str(start))


def add_page_field(paragraph, fmt, placeholder):
    def style(r):
        set_run(r, size=9, color=GRAY)
    r1 = paragraph.add_run()
    style(r1)
    f = OxmlElement("w:fldChar")
    f.set(qn("w:fldCharType"), "begin")
    r1._element.append(f)
    r2 = paragraph.add_run()
    style(r2)
    i = OxmlElement("w:instrText")
    i.set(qn("xml:space"), "preserve")
    i.text = " PAGE \\* %s \\* MERGEFORMAT " % fmt
    r2._element.append(i)
    r3 = paragraph.add_run()
    style(r3)
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "separate")
    r3._element.append(f3)
    r4 = paragraph.add_run(placeholder)
    style(r4)
    r5 = paragraph.add_run()
    style(r5)
    f5 = OxmlElement("w:fldChar")
    f5.set(qn("w:fldCharType"), "end")
    r5._element.append(f5)


def setup_section(sec, top=1440, bottom=1440, left=1701, right=1417):
    sec.page_width = Twips(11906)
    sec.page_height = Twips(16838)
    sec.top_margin = Twips(top)
    sec.bottom_margin = Twips(bottom)
    sec.left_margin = Twips(left)
    sec.right_margin = Twips(right)
    sec.header_distance = Twips(720)
    sec.footer_distance = Twips(720)


# ────────────────────────── 封面(R4:上深色块 + 分隔条 + 下白区) ──────────────────────────
def build_cover(doc):
    sec = doc.sections[0]
    setup_section(sec, top=0, bottom=0, left=0, right=0)

    t = doc.add_table(rows=3, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    table_width_pct(t)
    table_no_borders(t)
    table_cell_margins(t, top=0, bottom=0, left=1200, right=800)
    set_grid(t, [100])

    heights = [7500, 60, 9038]
    for i, h in enumerate(heights):
        trPr = t.rows[i]._tr.get_or_add_trPr()
        th = OxmlElement("w:trHeight")
        th.set(qn("w:val"), str(h))
        th.set(qn("w:hRule"), "exact")
        trPr.append(th)

    # 上部深色块
    c1 = t.cell(0, 0)
    cell_shading(c1, DARK)
    c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p0 = c1.paragraphs[0]
    pf = p0.paragraph_format
    pf.space_before = Twips(4100)
    pf.space_after = Twips(0)
    pf.line_spacing = 1.0

    pl = c1.add_paragraph()
    pl.paragraph_format.space_after = Twips(500)
    pl.paragraph_format.line_spacing = 1.0
    r = pl.add_run(" ".join("AI ENGINEERING WORKFLOW"))
    set_run(r, size=9, color=ACCENT, ascii_font="Calibri")

    pt_ = c1.add_paragraph()
    pf = pt_.paragraph_format
    pf.space_after = Twips(200)
    pf.line_spacing = Twips(920)
    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    r = pt_.add_run("AI 辅助工程学习方案")
    set_run(r, size=40, bold=True, color="FFFFFF", east_font="SimHei", ascii_font="Arial")

    ps = c1.add_paragraph()
    ps.paragraph_format.space_after = Twips(100)
    ps.paragraph_format.line_spacing = 1.0
    r = ps.add_run("STM32 测控系统毕业设计 · 人机分工与工程闭环实践")
    set_run(r, size=12, color="B0B8C0", ascii_font="Arial", east_font="Microsoft YaHei")

    # 强调色分隔条
    c2 = t.cell(1, 0)
    cell_shading(c2, ACCENT)
    p = c2.paragraphs[0]
    pf = p.paragraph_format
    pf.space_before = Twips(0)
    pf.space_after = Twips(0)
    pf.line_spacing = Twips(30)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pPr = p._p.get_or_add_pPr()
    rPrM = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    rPrM.append(sz)
    pPr.append(rPrM)

    # 下部白区
    c3 = t.cell(2, 0)
    c3.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p0 = c3.paragraphs[0]
    pf = p0.paragraph_format
    pf.space_before = Twips(800)
    pf.space_after = Twips(0)
    pf.line_spacing = 1.0
    meta = [
        "文档类型:个人学习与开发实施方案",
        "适用项目:基于 STM32F103C8T6 的测控系统毕业设计",
        "编 制 人:【姓名】",
        "编制日期:2026 年 9 月 13 日",
        "计划周期:16 周(2026 年 9 月 14 日 至 2027 年 1 月 3 日)",
    ]
    for m in meta:
        pm = c3.add_paragraph()
        pm.paragraph_format.space_after = Twips(120)
        pm.paragraph_format.line_spacing = 1.3
        r = pm.add_run(m)
        set_run(r, size=14, color=META, ascii_font="Arial", east_font="Microsoft YaHei")

    psp = c3.add_paragraph()
    psp.paragraph_format.space_before = Twips(1800)
    psp.paragraph_format.space_after = Twips(0)
    psp.paragraph_format.line_spacing = 1.0

    pft = c3.add_paragraph()
    pft.paragraph_format.line_spacing = 1.0
    r = pft.add_run("AI-ASSISTED ENGINEERING WORKFLOW")
    set_run(r, size=11, color="909090", ascii_font="Arial")
    r = pft.add_run("                                ")
    r = pft.add_run("2026.09")
    set_run(r, size=11, color="909090", ascii_font="Arial")

    # 表后收尾空段(极小行高,防止溢出到第 2 页)
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Twips(0)
    pf.space_after = Twips(0)
    pf.line_spacing = Twips(20)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pPr = p._p.get_or_add_pPr()
    rPrM = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    rPrM.append(sz)
    pPr.append(rPrM)


# ────────────────────────── 前置部分:执行摘要 + 目录 ──────────────────────────
ABSTRACT = [
    "本方案面向基于 STM32F103C8T6 的测控系统毕业设计课题,针对单人同时承担固件开发、上位机开发、硬件设计与测试验证所面临的精力与能力矛盾,提出一套人机分工、多模型协作的 AI 辅助工程学习方法。方案的核心是重新定位个人角色:本人承担项目经理与硬件工程师职责,负责需求定义、系统架构决策、硬件选型与接线、实物测试与最终验收;人工智能模型按工种组成虚拟团队,分别承担系统架构设计、快速原型编码、疑难故障诊断与代码审查工程化四类工作。",
    "方案规划了 16 周的四阶段实施路径:第 1 至 3 周补齐嵌入式基础与工具链,第 4 至 8 周完成 ADC、PWM、UART、超声波与 LabVIEW 上位机等功能模块的原型开发与实物验证,第 9 至 12 周完成系统集成、通信协议制定与 PCB 设计打样,第 13 至 16 周完成工程化交付与毕业设计材料整理。方案同时明确了四条流程红线:禁止\u201cAI 写代码、烧录、能跑即完事\u201d的工作模式;所有 AI 产出的结论必须经过示波器、万用表等实测手段验证;每个关键设计点背后的\u201c为什么\u201d必须能够独立回答;实物安全优先于进度。",
    "按本方案执行,最终交付物不仅是一套可运行的测控系统(硬件原理图与 PCB、STM32 固件工程、LabVIEW 上位机、测试记录与文档),更是一套可复用的 AI 辅助工程工作流,以及一段面向质量、工艺、项目与自动化类岗位、有完整过程记录支撑的项目履历。方案末章还基于公开技术社区资料设计了创新点与答辩准备要点,并整理了分类的参考资源清单,便于各阶段按图索骥。",
]


def build_front_matter(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    setup_section(sec)
    set_pgnum(sec, "upperRoman", 1)
    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp, "ROMAN", "I")

    heading(doc, "执行摘要", 1)
    for i, t in enumerate(ABSTRACT):
        p = body(doc, t)
        if i == len(ABSTRACT) - 1:
            p.add_run().add_break(WD_BREAK.PAGE)

    # 目录标题:不使用 Heading 样式,防止目录索引自身
    para(doc, "目  录", size=16, bold=True, east_font="SimHei",
         align=WD_ALIGN_PARAGRAPH.CENTER, before=240, after=240)

    # TOC 域
    p = doc.add_paragraph()
    r1 = p.add_run()
    f = OxmlElement("w:fldChar")
    f.set(qn("w:fldCharType"), "begin")
    f.set(qn("w:dirty"), "true")
    r1._element.append(f)
    r2 = p.add_run()
    i = OxmlElement("w:instrText")
    i.set(qn("xml:space"), "preserve")
    i.text = ' TOC \\o "1-3" \\h \\z \\u '
    r2._element.append(i)
    r3 = p.add_run()
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "separate")
    r3._element.append(f3)
    # end 域字符必须独立成段:占位条目将被注入 separate 段与 end 段之间(域内),
    # 否则条目落在域外,Word 更新域后会出现两份目录
    p2 = doc.add_paragraph()
    p2.paragraph_format.line_spacing = 1.3
    r4 = p2.add_run()
    f4 = OxmlElement("w:fldChar")
    f4.set(qn("w:fldCharType"), "end")
    r4._element.append(f4)

    para(doc, "注:本目录由域代码自动生成。编辑文档后请在目录上单击鼠标右键,选择\u201c更新域→更新整个目录\u201d,以确保页码准确。",
         size=9, italic=True, color="888888", before=200)


# ────────────────────────── 正文内容 ──────────────────────────
def build_body(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    setup_section(sec)
    set_pgnum(sec, "decimal", 1)
    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp, "arabic", "1")
    sec.header.is_linked_to_previous = False
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run("AI 辅助工程学习方案")
    set_run(r, size=9, color=GRAY)

    ch1(doc)
    ch2(doc)
    ch3(doc)
    ch4(doc)
    ch5(doc)
    ch6(doc)
    ch7(doc)
    ch8(doc)
    ch9(doc)
    ch10(doc)
    ch11(doc)


def ch1(doc):
    heading(doc, "1 方案背景与总体目标", 1)
    heading(doc, "1.1 背景与现状分析", 2)
    body(doc, "毕业设计要求在有限周期内独立完成 STM32F103C8T6 测控系统的全部环节,涵盖 C 语言固件开发、LabVIEW 上位机设计、PCB 设计、通信协议制定与系统测试。该工作量覆盖了传统上需要一个小型团队分工完成的多种工种,若全部依靠个人逐行手写代码、逐项查阅资料,进度与质量都难以同时保证。")
    body(doc, "与此同时,以大语言模型为代表的 AI 工具已经能够较好地完成代码生成、方案设计、故障分析与文档撰写等工作,且不同规格的模型在成本与能力上差异明显:轻量模型响应快、调用成本低,适合大量试错与快速迭代;高规格模型推理能力强,适合疑难问题的深度分析。如果不对模型加以组织、混用各类模型,又容易出现产出风格不一、质量不可控、带缺陷的代码直接进入实物环节等问题。")
    body(doc, "本方案的基本思路是:不追求亲自编写全部代码,而是把个人定位为系统架构师与项目经理,让不同模型按工种分工协作,并以\u201c规格先行、实物验证\u201d的工程闭环约束 AI 的产出质量,使其真正服务于毕业设计目标。")
    body(doc, "从评价体系看,各高校公开的毕业设计管理办法普遍采用\u201c指导教师+评阅教师+答辩\u201d三部分加权评分(常见 3:3:4),创新性是评定优秀的关键条件;公开的选题经验同时指出,单纯\u201c传感器采集+显示\u201d类课题同质化严重,难以体现工作量与创新能力。因此,本方案在保证功能实现之外,第十章专门设计了创新点与答辩准备,第十一章给出经检索整理的参考资源清单,使工作量与创新性都有据可依。")

    heading(doc, "1.2 总体目标", 2)
    body(doc, "目标一:完成毕业设计实物系统。包括基于 STM32F103C8T6 的测控板,实现电压与电流的 ADC 采集、舵机与直流电机的 PWM 控制、UART 通信与超声波测距;配套 LabVIEW 上位机实现监控与数据采集;系统整体联调通过并稳定运行。")
    body(doc, "目标二:建立并跑通 AI 辅助工程闭环。形成\u201c需求规格→快速实现→实物测试→深度诊断→代码审查→回归测试→交付\u201d的标准流程,并将该流程沉淀为个人可复用的工程方法。")
    body(doc, "目标三:形成完整的工程资产。按统一的目录结构管理需求规格、系统架构、硬件设计、固件工程、上位机、测试记录与交付文档,达到可直接用于毕业答辩与求职展示的完整度。")

    heading(doc, "1.3 角色定位与职责边界", 2)
    body(doc, "在本方案中,个人角色为项目经理兼硬件工程师。以下职责不可外包给 AI:最终需求定义与取舍;系统架构与技术路线决策;元器件选型与 PCB 确认;实物接线与供电设计的确认;示波器与万用表实测;故障现象的初步定位与复现;需求是否完成的最终判断与验收。")
    body(doc, "尤其对于电源、电机、传感器等涉及实物安全的部分,AI 给出的\u201c应该可以工作\u201d与示波器上真实测到波形是两回事。本方案约定:一切技术结论以实测为准,AI 的任何判断仅作为待验证的建议。")


def ch2(doc):
    heading(doc, "2 AI 工程团队的分工与职责边界", 1)
    heading(doc, "2.1 分工总则", 2)
    body(doc, "AI 团队按工种而非模型品牌划分。每个工种选择该类任务性价比最高的模型;所有模型的产出一律视为\u201c建议\u201d,须经本人确认或实测通过后方可进入下一环节。各角色的职责与使用要点如表 1 所示。")
    add_data_table(
        doc, "表 1 AI 工程团队角色分工", [19, 23, 30, 28],
        ["角色", "承担者", "核心职责", "使用要点"],
        [
            ["项目经理 /\n硬件工程师", "本人",
             "需求定义、系统架构决策、硬件选型与接线、实物测试、进度控制、最终验收",
             "唯一不可外包的角色;\n对进度与质量\n负最终责任"],
            ["系统架构师 /\n技术顾问", "ChatGPT 等\n对话强模型",
             "总体架构、技术路线、模块接口、通信协议、测试方案与文档框架设计",
             "在动手编码前先解决\n\u201c怎么做\u201d;输出方案\n文档而非代码"],
            ["快速原型\n工程师", "DeepSeek /\nGLM 等轻量模型",
             "各功能模块 demo 代码的快速生成与迭代(如 uart.c、adc.c、servo.c)",
             "允许大量试错;先求能跑再谈质量;单次任务只做一个模块"],
            ["疑难问题\n诊断专家", "高规格模型\n(DeepSeek Pro、\nGLM 高规格等)",
             "数据漂移、舵机抖动、\n丢包、乱码等疑难\n故障的根因分析",
             "提交完整证据包:\n代码+原理图+\n串口日志+波形截图\n+现象描述"],
            ["代码审查\n与工程化", "Codex",
             "代码审查、重构、\n测试补充、工程结构\n规范化、最终交付检查",
             "在原型可运行后介入;\n交付前与最终交付\n各执行一次"],
        ])
    body(doc, "需要强调的是,诊断专家与架构师两类角色引入高成本模型,目的是把昂贵模型的推理能力用在刀刃上;而日常代码生成与修改交给轻量模型,利用其低成本优势换取足够多的迭代次数。")

    heading(doc, "2.2 模型调用策略", 2)
    body(doc, "按成本与能力匹配任务:轻量模型承担约八成的常规编码与迭代任务,可反复生成与修改而不必顾虑成本;高规格模型仅用于疑难诊断与关键方案评审;Codex 承担审查与工程化。每次调用都应提供完整上下文,包括需求说明、相关代码、实测现象、串口日志与波形截图,避免模型凭空猜测。对同一疑难问题,可交叉询问两个不同模型并比对结论,以降低单一模型误判的风险。")
    body(doc, "这一分工与业界当前的主流实践方向一致:公开的 LLM 编码工作流综述(如 Addy Osmani 在 2026 年初发表的工作流总结)普遍主张\u201c先与 AI 共同打磨详细规格、再形成分步计划、最后进入实现\u201d的规格先行模式;嵌入式领域的头部企业也在推行\u201c一切皆代码(Everything as Code)\u201d,使 AI 能够参与开发全流程。相应地,Codex、Claude Code 等代理型编码工具在本方案中的定位是编码与审查的加速器,而不是技术决策者。")


def ch3(doc):
    heading(doc, "3 工程闭环工作流程", 1)
    heading(doc, "3.1 流程总览", 2)
    body(doc, "标准流程共七个环节,构成闭环:第一,需求规格,由本人撰写模块级需求(输入、输出、精度、时序),必要时请架构师模型评审;第二,快速实现,轻量模型依据规格生成 demo 代码;第三,实物测试,烧录后用仪器验证功能;第四,深度诊断,未通过时由高规格模型基于证据包分析根因;第五,代码审查,Codex 对通过测试的代码进行审查与重构;第六,回归测试,重构后重新实物验证;第七,交付,Codex 执行最终工程化检查并归档。整个流程中,本人始终处于决策与验收位置,模型在环节内部承担具体工作。该闭环与业界规格先行的 LLM 工作流共识一致,并以实物测试作为 AI 产出的最终裁判。")
    body(doc, "各环节的执行者、输入输出与通过标准如表 2 所示。")
    add_data_table(
        doc, "表 2 工程闭环各环节定义与通过标准", [13, 15, 20, 24, 28],
        ["环节", "执行者", "输入", "输出", "通过标准"],
        [
            ["① 需求规格", "本人\n(架构师\n模型辅助)", "系统需求", "模块需求规格书",
             "接口、参数、验收方法齐备"],
            ["② 快速实现", "轻量模型", "需求规格书", "模块 demo 代码",
             "编译通过,逻辑与规格一致"],
            ["③ 实物测试", "本人", "demo 代码+硬件", "测试记录",
             "功能达标且现象可复现"],
            ["④ 深度诊断", "高规格模型", "证据包", "根因分析与修改建议",
             "建议经实测确认有效"],
            ["⑤ 代码审查", "Codex", "通过测试的代码", "审查意见+重构代码",
             "无缓冲区、超时、中断冲突等明显隐患"],
            ["⑥ 回归测试", "本人", "重构后代码", "回归测试记录",
             "原有功能全部保持"],
            ["⑦ 交付", "Codex+本人", "全部工程文件", "交付工程包",
             "目录、文档、版本齐备"],
        ], center_cols=(0,))

    heading(doc, "3.2 流程红线", 2)
    body(doc, "红线一:禁止\u201cAI 写代码→烧录→能跑即完事\u201d。若长期停留在该模式,个人将退化为\u201cAI 操作员\u201d而不是硬件工程师,答辩与面试时无法解释系统的工作原理,项目也不具备可信度。")
    body(doc, "红线二:AI 结论必须实测验证。无论模型给出多么确定的判断,进入硬件环节前都必须经过万用表、示波器或串口日志的实测确认。")
    body(doc, "红线三:每个\u201c为什么\u201d必须能回答。本方案第五章列出了必须亲自掌握的核心问题清单,对其中任一问题若不能解释原理并通过实验验证,对应模块不得进入交付状态。")
    body(doc, "红线四:实物安全优先。任何新硬件组合在首次上电前,必须完成电源极性、共地关系与限流措施检查;电机、舵机等大电流负载必须独立供电。")


def ch4(doc):
    heading(doc, "4 分阶段学习与实施计划(16 周)", 1)
    heading(doc, "4.1 阶段总览", 2)
    body(doc, "计划自 2026 年 9 月 14 日(第 1 周周一)起,共 16 周,分四个阶段推进。每个阶段设置明确的关键成果,作为周检查与里程碑验收的依据,详见表 3。")
    add_data_table(
        doc, "表 3 十六周阶段总览", [12, 26, 24, 38],
        ["阶段", "周期", "主题", "关键成果"],
        [
            ["阶段一", "第 1–3 周\n9 月 14 日–10 月 4 日", "嵌入式基础\n与工具链",
             "开发环境跑通;GPIO、定时器、中断、串口四个基础实验实物通过"],
            ["阶段二", "第 4–8 周\n10 月 5 日–11 月 8 日", "功能模块开发\n与实物验证",
             "ADC、PWM、UART、超声波、LabVIEW 五个模块逐一通过实测"],
            ["阶段三", "第 9–12 周\n11 月 9 日–12 月 6 日", "系统集成\n与 PCB",
             "通信协议定稿;系统联调通过;PCB 完成打样并验证"],
            ["阶段四", "第 13–16 周\n12 月 7 日–2027 年 1 月 3 日", "工程化交付\n与毕设材料",
             "交付工程包、测试报告与毕业论文初稿"],
        ], center_cols=(0,))

    heading(doc, "4.2 阶段一:嵌入式基础与工具链(第 1–3 周)", 2)
    body(doc, "本阶段的目标是建立最小系统开发能力,为后续模块开发打底。学习以\u201c概念理解+实验验证\u201d的方式进行,所有示例代码都要求能够脱稿复述其原理。实施要点见表 4。")
    phase_table(doc, "表 4 阶段一实施要点", [
        ["学习目标", "掌握 STM32 最小系统开发能力,理解时钟、GPIO、定时器与中断的基本机制;以 STM32CubeMX + HAL 库为主线搭建工程(标准库仅作对照阅读),熟悉 Keil 编译烧录与调试方法。"],
        ["学习内容", "C 语言指针、结构体与位操作回顾\nSTM32F103C8T6 最小系统与时钟树\n(参考第十一章设计文章与视频)\nGPIO 输入输出与按键消抖\n定时器与 PWM 基本概念\n中断与 NVIC 基础\nUART 串口通信基础\nSTM32CubeMX + HAL 库工程建立流程"],
        ["AI 协作方式", "向架构师模型索取学习路线与关键概念讲解;用轻量模型生成 GPIO 点灯、定时器中断等示例代码;每个示例必须能向模型复述工作原理并请其纠错;教程以第十一章清单中的系列课程与文章为主线。"],
        ["实践任务", "搭建最小系统板;完成 LED 点灯、按键输入、定时器中断、串口收发四个基础实验。"],
        ["交付物", "基础实验代码仓库、学习笔记、开发环境说明。"],
        ["验收标准", "四个基础实验全部实物通过;能不借助资料画出最小系统框图并解释时钟分配关系。"],
    ])

    heading(doc, "4.3 阶段二:功能模块开发与实物验证(第 4–8 周)", 2)
    body(doc, "本阶段进入核心功能开发,每周完成一个模块:电压电流采集、舵机与电机控制、UART 通信、超声波测距、LabVIEW 上位机。每个模块均走完整的工程闭环,实施要点见表 5。")
    phase_table(doc, "表 5 阶段二实施要点", [
        ["学习目标", "完成 ADC(电压/电流)、PWM(舵机/直流电机)、UART 协议、超声波测距与 LabVIEW 上位机五大模块的原型开发,逐一通过实物验证,并积累实测数据。"],
        ["学习内容", "ADC 采样原理、参考电压与软件滤波\n舵机 50Hz 控制原理与电机驱动电路(MOSFET 栅极电阻、独立供电、共地)\n电机驱动选型对比(TB6612 与 L298N 的压降与效率差异)\n电流检测双方案:采样电阻+运放接 ADC,与 INA219(I2C)方案对比\nUART 帧格式、波特率与接收缓冲区\n超声波测距时序与超时处理\nLabVIEW VISA 串口编程与前面板设计"],
        ["AI 协作方式", "本人编写模块需求规格;轻量模型生成 demo 并快速迭代;示波器与万用表逐项实测;异常现象整理证据包后交高规格模型诊断。"],
        ["实践任务", "五个模块按周推进;每模块留存测试照片、波形截图与串口日志。"],
        ["交付物", "各模块可运行代码、模块测试记录、问题诊断记录。"],
        ["验收标准", "每个模块有实测数据(如电压测量误差、舵机角度精度、连续通信丢包率)并达到自定指标;核心问题清单中对应条目可独立回答。"],
    ])

    heading(doc, "4.4 阶段三:系统集成与 PCB(第 9–12 周)", 2)
    body(doc, "本阶段把各模块集成为完整系统:制定并冻结上下位机通信协议,开展系统联调,同时完成 PCB 设计与打样,为最终交付形态做准备。实施要点见表 6。")
    phase_table(doc, "表 6 阶段三实施要点", [
        ["学习目标", "将各模块集成为完整系统;制定并冻结通信协议;完成系统联调与 PCB 设计打样。"],
        ["学习内容", "帧头、长度、校验和的协议设计方法\n主循环与中断服务的分工\n系统级供电与去耦设计\n原理图审查要点\nPCB 布局布线基础(电源路径、晶振、模拟数字分区)"],
        ["AI 协作方式", "通信协议由架构师模型参与设计、本人定稿;Codex 对整合后的代码做第一次全面审查与重构;原理图与 PCB 关键网络请架构师模型按检查清单复核,本人最终确认。"],
        ["实践任务", "协议联调(含丢包与校验测试);整机长时间可靠性预运行;PCB 布局与打样、焊接、上板验证。"],
        ["交付物", "通信协议文档、系统联调记录、PCB 工程文件与打样板。"],
        ["验收标准", "上位机与下位机按协议稳定通信;PCB 板焊接后功能与原型版本一致。"],
    ])

    heading(doc, "4.5 阶段四:工程化交付与毕设材料(第 13–16 周)", 2)
    body(doc, "本阶段完成代码工程化、测试收尾、文档与论文撰写,把开发成果整理为可交付、可答辩、可求职展示的完整资产。实施要点见表 7。")
    phase_table(doc, "表 7 阶段四实施要点", [
        ["学习目标", "完成代码工程化、全量回归测试、文档与论文撰写,形成完整可交付资产。"],
        ["学习内容", "代码规范与模块化整理\n测试报告撰写方法\n论文结构与图表规范\n答辩要点提炼"],
        ["AI 协作方式", "Codex 执行最终代码审查、补全注释与 README;强对话模型协助搭建论文框架与文字润色;技术内容与数据全部由本人提供并逐项核对。"],
        ["实践任务", "全量回归测试并整理测试记录;撰写需求规格书、系统架构文档、硬件测试报告、用户手册;完成论文初稿与答辩材料。"],
        ["交付物", "按第七章目录结构组织的最终交付工程包、论文初稿、答辩演示材料。"],
        ["验收标准", "交付包可在他人电脑上按 README 复现编译与烧录;论文中引用的数据与测试记录一一对应。"],
    ])


def ch5(doc):
    heading(doc, "5 必须亲自掌握的核心知识清单", 1)
    heading(doc, "5.1 清单使用说明", 2)
    body(doc, "本章清单覆盖系统各关键设计点,是防止\u201c只会用 AI、不懂原理\u201d的底线机制。处理方式为三步:先请 AI 把原理讲透,再脱稿复述并让 AI 挑错,最后设计小实验加以证实。全部 13 项须在阶段四开始前完成;对应模块在所属问题通过验证前,不得标记为完成。")
    heading(doc, "5.2 核心问题与验证方式", 2)
    add_data_table(
        doc, "表 8 必须亲自掌握的核心问题清单", [8, 40, 16, 36],
        ["编号", "核心问题", "知识领域", "验证方式"],
        [
            ["Q1", "为什么本系统需要 ADC?采样值如何换算成真实电压与电流?", "模拟量采集", "用已知电压源与分压电阻比对换算结果"],
            ["Q2", "为什么舵机控制信号是 50Hz、0.5–2.5ms 脉宽?", "PWM 应用", "示波器测量 PWM 周期与脉宽,改变占空比观察舵机角度"],
            ["Q3", "为什么 MOSFET 栅极要串联电阻?", "功率驱动", "对比加与不加电阻时的驱动波形振铃与器件发热"],
            ["Q4", "为什么电机必须独立供电?", "电源设计", "测量电机启动瞬间\n主电源电压跌落"],
            ["Q5", "为什么 STM32 与外部\n设备必须共地?", "电源设计", "断开共地,观察串口乱码与电平异常"],
            ["Q6", "为什么 UART 在 115200 波特率下会出现乱码?", "串口通信", "示波器实测波特率并与时钟误差比对"],
            ["Q7", "为什么 UART 会丢包?如何从协议层面解决?", "串口通信", "高速连续发送统计丢包率,加入校验与重传后复测"],
            ["Q8", "为什么 ADC 采样会有噪声?有哪些抑制手段?", "信号处理", "对比软件均值滤波\n前后的数据波动"],
            ["Q9", "为什么电源纹波会影响传感器精度?", "电源完整性", "观察电机运行与停止时 ADC 读数的变化"],
            ["Q10", "为什么定时器 PWM 分辨率与预分频有关?", "定时器", "计算并实测不同预分频下的最小脉宽步进"],
            ["Q11", "超声波测距的时序要求是什么?为什么需要超时判断?", "传感器", "人为遮挡与超量程测试,观察程序行为"],
            ["Q12", "LabVIEW VISA 读数据为什么会阻塞?如何设计缓冲与超时?", "上位机", "分别测试固定长度读取与带超时读取的表现"],
            ["Q13", "为什么低压小电机驱动优先选 TB6612,而不是 L298N?", "功率驱动", "相同 PWM 占空比下对比电机端电压与驱动板发热"],
        ], center_cols=(0,))


def ch6(doc):
    heading(doc, "6 测试与实物验证方法", 1)
    heading(doc, "6.1 测试原则", 2)
    body(doc, "测试遵循\u201c先模块、后系统;先静态、后上电;先开环、后闭环\u201d的顺序。任何模块在通过实物测试前不得宣称完成;所有测试结论以仪器读数与记录为准,不以系统\u201c看起来运行正常\u201d为准。")
    heading(doc, "6.2 测试工具与典型检查项", 2)
    body(doc, "各测试工具的用途与典型检查项如表 9 所示,其中示波器是本方案中区分\u201cAI 认为正常\u201d与\u201c实际正常\u201d的关键工具。")
    add_data_table(
        doc, "表 9 测试工具与典型检查项", [20, 32, 48],
        ["工具", "主要用途", "典型检查项"],
        [
            ["万用表", "电压、电流、通断测量", "上电前电源短路检查;各节点电平;整机与模块工作电流"],
            ["示波器", "波形与时序测量", "PWM 周期与脉宽;驱动波形振铃;UART 波形;电源纹波;ADC 输入端噪声"],
            ["USB 转 TTL+串口助手", "调试信息输出与协议观测", "中间变量打印;帧格式核对;连续收发丢包统计"],
            ["LabVIEW 上位机", "系统联调与数据记录", "曲线显示正确性;数据保存完整性;异常值捕捉"],
        ])
    heading(doc, "6.3 测试记录规范", 2)
    body(doc, "每条测试记录包含八个字段:测试编号、日期、被测对象、测试条件(供电电压、负载状态)、操作步骤、预期结果、实测结果、结论与后续动作。测试记录集中存放于项目 05_Test 目录,并与 Bug 记录联动:凡实测结果与预期不符,立即登记 Bug 并按第三章流程进入诊断环节。")


def ch7(doc):
    heading(doc, "7 工程资产与文档管理", 1)
    heading(doc, "7.1 项目目录结构", 2)
    body(doc, "全部工程资产按统一的目录结构管理,使 AI 在每次任务中都能获得完整上下文,也使交付物可直接用于答辩与求职展示。目录结构如表 10 所示。")
    add_data_table(
        doc, "表 10 项目工程资产目录结构", [26, 74],
        ["目录", "存放内容"],
        [
            ["00_Requirements", "需求规格书、需求变更记录"],
            ["01_System", "系统架构说明、通信协议文档"],
            ["02_Hardware", "原理图、PCB 工程、BOM 清单、硬件测试报告"],
            ["03_Firmware", "STM32 工程、各模块代码、README 与编译烧录说明"],
            ["04_LabVIEW", "上位机 VI 工程与打包程序"],
            ["05_Test", "测试方案、测试记录、Bug 记录、波形截图与串口日志"],
            ["06_Delivery", "用户手册、技术文档、答辩与毕业设计材料"],
        ])
    heading(doc, "7.2 版本与文档规范", 2)
    body(doc, "所有代码与文档纳入 Git 管理:main 分支始终保持可交付状态,功能开发在独立分支进行,经 Codex 审查并实测通过后方可合并。需求发生变更时,须同步更新需求规格书,并在变更记录中注明日期与原因。AI 生成的代码入库前必须通过实物回归测试;每次重要调试的示波器截图与串口日志按测试编号归档,作为后续诊断证据包的素材。")


def ch8(doc):
    heading(doc, "8 风险分析与应对措施", 1)
    body(doc, "结合单人开发与 AI 协作的特点,识别出六类主要风险,如表 11 所示。风险应对的基本原则是:硬件备份与限流上电保安全,回归测试与双阶段审查保质量,里程碑核对保进度。")
    add_data_table(
        doc, "表 11 主要风险与应对措施", [22, 14, 64],
        ["风险", "影响程度", "应对措施"],
        [
            ["硬件损坏\n导致进度中断", "高",
             "关键元器件(最小系统板、驱动模块)准备备份;首次上电使用限流电源;先测电源、后接负载"],
            ["AI 代码\n存在隐蔽缺陷", "高",
             "坚持实物回归测试;Codex 双阶段审查;将缓冲区、超时、中断冲突列为固定检查项"],
            ["学习进度\n滞后于计划", "中",
             "每周末核对阶段里程碑;允许调整模块开发顺序但不降低验收标准;必要时压缩 PCB 打样周期"],
            ["强模型\n诊断结论误导", "中",
             "诊断结论必须设计验证实验确认;同一问题交叉询问两个模型并比对结论"],
            ["文档与代码\n不同步", "低",
             "文档更新纳入各阶段交付物检查;代码合并时同步检查 README 与协议文档"],
            ["SWD 调试\n连接不稳定", "中",
             "按社区经验排查 SWD 接线、复位电路与驱动安装;保留串口 ISP 与 BOOT 跳线作为备用烧录通道"],
        ], center_cols=(1,))


def ch9(doc):
    heading(doc, "9 预期成果与能力沉淀", 1)
    heading(doc, "9.1 预期交付成果", 2)
    body(doc, "完成本方案后,应交付:一套通过联调的测控系统(含 PCB 实物)、STM32 固件工程、LabVIEW 上位机、完整测试记录与硬件测试报告,以及按统一目录组织的工程文档。其中,各阶段积累的实测数据与诊断记录构成毕业论文的核心素材。")
    heading(doc, "9.2 能力沉淀与求职表述", 2)
    body(doc, "本方案真正沉淀的能力包括:系统需求定义与拆解、硬件架构与选型判断、以实测为准的验证习惯、多模型协作的组织方法,以及对通信协议、电源设计、信号采集等原理的亲身理解。这些能力可以通过一段完整的经历来表述:")
    p = para(doc, "\u201c本人负责系统需求与硬件架构,设计了基于 STM32F103C8T6 的测控板,制定了上下位机通信协议,采用 AI 辅助方式完成固件开发,通过示波器与万用表完成硬件验证,使用 LabVIEW 实现上位机监控与数据采集,最终完成系统联调与交付。\u201d",
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=480, after=80)
    p.runs[0].font.bold = True
    body(doc, "这一表述背后有完整的过程记录与工程资产支撑,其含金量显著高于单纯的\u201c会使用 STM32 和 LabVIEW\u201d。项目结束后,本方案所定义的工作流本身也可以作为方法论沉淀,迁移到后续任何软硬件结合的开发任务中。")


def ch10(doc):
    heading(doc, "10 创新点设计与答辩准备", 1)
    heading(doc, "10.1 创新点设计", 2)
    body(doc, "根据公开的毕业设计评分标准与选题经验,创新性是评定优秀的关键条件,而单纯\u201c采集+显示\u201d类课题难以支撑创新性评价。结合本课题的测控定位,从检索到的主流创新方向(控制算法优化、上位机数据管理、通信可靠性设计、物联扩展)中筛选出以下候选创新点,并给出落地方式与取舍建议,见表 12。")
    add_data_table(
        doc, "表 12 创新点设计与取舍建议", [20, 36, 14, 30],
        ["创新点", "在本课题中的落地方式", "实现成本", "建议"],
        [
            ["直流电机 PID\n调速闭环", "TB6612 驱动带编码器电机,增量式 PID 闭环调速,LabVIEW 实时显示设定与实际转速曲线", "中", "推荐采纳\n阶段三实现"],
            ["上位机数据管理", "LabVIEW 历史数据存储与导出、曲线回放与异常区间标注", "低", "推荐采纳\n阶段二打底"],
            ["通信可靠性设计", "帧头+长度+校验和\n+超时重传的完整协议,\n并给出丢包率对比数据", "低", "推荐采纳\n已纳入阶段三"],
            ["双方案\n电流测量对比", "采样电阻+运放接 ADC 与 INA219(I2C)两种方案的实测精度对比,形成论文数据", "中", "可选\n视阶段二进度"],
            ["物联远程监控", "ESP8266 接入云平台,手机端查看测控数据", "高", "备选\n进度充裕时启动"],
            ["AI 辅助\n开发工作流", "将本方案的人机分工与工程闭环作为方法论写入论文,以测试记录佐证其有效性", "低", "推荐采纳\n贯穿全程"],
        ], center_cols=(2,))
    heading(doc, "10.2 答辩准备要点", 2)
    body(doc, "公开的高校答辩要求显示:答辩一般占毕业设计总评的四成左右,通常要求 10 分钟左右的陈述 PPT,鼓励携带笔记本电脑进行现场实物演示;论文正文普遍要求不少于 1.5 万字,格式不规范可能直接影响成绩评定。据此,阶段四须落实以下准备:制作 10 分钟陈述 PPT 并录制演示视频作为备份;准备可现场运行的实物与上位机;按学校模板核对论文格式与图表规范;组织不少于两次模拟答辩;预先准备评委高频问题的回答,包括选题依据、创新点、关键参数取舍(如 50Hz 舵机信号、电机驱动选型)与实测数据来源。")


def ch11(doc):
    heading(doc, "11 参考资源清单", 1)
    heading(doc, "11.1 使用说明", 2)
    body(doc, "以下资源于 2026 年 9 月 13 日经搜索引擎检索整理,按官方与社区文档、CSDN 实战与避坑文章、GitHub 开源项目、视频课程与选题答辩参考四类组织。资源定位为对照学习与避坑参考,不能替代本人的实物测试与验证;引用进论文前须自行核对原文与版本时效。")
    heading(doc, "11.2 官方与社区文档", 2)
    body(doc, "官方与社区文档的权威性最高,优先用于确认 API 用法与硬件原理,见表 13。")
    add_data_table(
        doc, "表 13 官方与社区文档", [56, 44],
        ["资源", "说明与建议用途"],
        [
            ["STM32 Serial Port Control with LabVIEW\n(NI Community 论坛)\nforums.ni.com/t5/LabVIEW-Embedded/STM32-Serial-Port-Control-with-LabView/td-p/4091331",
             "LabVIEW 用 VISA 配置串口与 STM32 通信的官方论坛帖;阶段二上位机开发参考"],
            ["ST 社区:Nucleo 与 LabVIEW 串口通信讨论\ncommunity.st.com/others-stm32-mcus-related-142/regarding-serial-communication-between-stm32-nucleo-f401re-and-labview-145428",
             "上位机向 STM32 下发控制命令的双向通信讨论;阶段三协议设计参考"],
            ["野火《电机应用开发实战指南》电流电压采集章\ndoc.embedfire.com/motor/motor_tutorial/zh/latest/basis_part/brushed_current_vbus_gather.html",
             "母线电流采样的采样电阻+运放+ADC 架构讲解;阶段二电流检测主线方案"],
            ["TI 电流检测芯片中文资料(含 INA219 系列)\nti.com.cn/cn/lit/pdf/zhcac20",
             "TI 电流检测官方资料与 INA232 迁移说明;电流方案精度升级时的进阶阅读"],
            ["SimpleFOC 低侧电流采样文档(中文)\ndocs.simplefoc.com/docs_chinese/low_side_current_sense",
             "分流电阻+增益+ADC 的低侧采样架构;理解电流采样原理的进阶材料"],
        ], size=9)
    heading(doc, "11.3 CSDN 实战与避坑文章", 2)
    body(doc, "CSDN 上同平台(C8T6)、同课题类型的实战文章最多,适合作为各阶段的对照材料,见表 14。")
    add_data_table(
        doc, "表 14 CSDN 实战与避坑文章", [56, 44],
        ["资源", "说明与建议用途"],
        [
            ["基于 STM32F103C8T6 的毕业设计入门实战:从外设配置到系统集成\nblog.csdn.net/2600_94959951/article/details/159419574",
             "外设配置到系统集成的完整路线,含 SWD 调试不稳定等高频坑点;阶段一通读"],
            ["STM32 毕业设计选题避坑指南\nblog.csdn.net/weixin_36483050/article/details/157951626",
             "从芯片资源、外设复杂度、可观测性评估选题门槛;已定题者用作风险自查"],
            ["基于 STM32F103C8T6 的智能环境监测系统设计与实现\nblog.csdn.net/weixin_33910191/article/details/164607333",
             "同类测控项目的选型思路与系统架构参考,含芯片选型理由分析"],
            ["STM32F103C8T6 最小系统设计(电源与 PCB)\nblog.csdn.net/weixin_45844053/article/details/136843955",
             "5V 转 3.3V 电源实用技巧与 PCB 打板注意事项;阶段三硬件设计前必读"],
            ["STM32F103C8T6 最小系统设计全指南(工业级)\nblog.csdn.net/rust6ferris/article/details/155739890",
             "电源、晶振、复位、BOOT 四大子系统可靠性设计;PCB 阶段对照检查清单"],
            ["HAL 库 + STM32CubeMX 系列学习教程\nblog.csdn.net/as480133937/article/details/99935090",
             "CubeMX 图形化配置与 HAL 库入门主线教程;阶段一使用"],
            ["STM32 开发:INA219 采集电压、电流值\nblog.csdn.net/weixin_48991062/article/details/134292801",
             "INA219 I2C 电流检测教程;\n阶段二对比方案"],
            ["STM32 控制 TT 马达:从 H 桥原理到 TB6612FNG 实战\nbbs.csdn.net/weixin_29181861/article/details/100282907",
             "L298N 压降问题的量化分析与 TB6612 实战;阶段二电机驱动参考"],
        ], size=9)
    heading(doc, "11.4 开源项目参考", 2)
    body(doc, "开源项目用于对标上位机形态与数据组织方式,重点学习其帧解析与界面设计思路,而不是直接照搬,见表 15。")
    add_data_table(
        doc, "表 15 开源项目参考", [56, 44],
        ["资源", "说明与建议用途"],
        [
            ["Serial-Studio\ngithub.com/Serial-Studio/Serial-Studio",
             "开源串口遥测仪表盘,支持 STM32 等设备接入;上位机界面与数据组织思路的对标对象"],
            ["xiahouzuoxin/zxSerial\ngithub.com/xiahouzuoxin/zxSerial",
             "基于 LabVIEW 的串口波形上位机;波形显示与帧解析的实现参考"],
            ["xiahouzuoxin/notes 中的《LabVIEW 上位机与串口通信》笔记\ngithub.com/xiahouzuoxin/notes\n(essays 目录)",
             "帧解析、多通道数据与带时间戳存储的文字教程;阶段二阅读"],
            ["HoussemLahmar/STM32_data_logger\ngithub.com/HoussemLahmar/STM32_data_logger",
             "STM32 模拟量到 PC 的数据记录项目;上位机数据落盘格式的参考"],
        ], size=9)
    heading(doc, "11.5 视频课程与选题、答辩参考", 2)
    body(doc, "视频课程适合入门建立直觉,选题与答辩资料用于对照评分要求校准工作量,见表 16。")
    add_data_table(
        doc, "表 16 视频课程与选题、答辩参考", [56, 44],
        ["资源", "说明与建议用途"],
        [
            ["《STM32F103C8T6 从入门到精通》\n最小系统原理图解读(B 站)\nbilibili.com/video/BV1Ui4y1W7SY",
             "零基础最小系统原理图视频讲解;阶段一开板前观看"],
            ["labview + STM32\n温湿度数据采集演示(B 站)\nbilibili.com/video/BV1fA411Q77A",
             "STM32 与 LabVIEW 上位机联动的完整效果演示;阶段二交付形态参考"],
            ["江协科技、正点原子、\n韦东山 STM32 系列课程(B 站)\nbilibili.com(搜索课程名)",
             "口碑较好的系统课程;江协科技适合入门打基础,正点原子配套资料齐全"],
            ["《STM32 初学者,到底选\n标准库还是 HAL 库?》(知乎)\nzhuanlan.zhihu.com/p/648673431",
             "标准库与 HAL 库的取舍分析;阶段一路线决策参考"],
            ["西北工业大学生命学院毕业设计成绩评定标准\nshengming.nwpu.edu.cn/info/1033/4248.htm",
             "公开的评分标准样例;对照理解创新性与规范性要求"],
            ["STM32 毕业设计题目 350 例(CSDN)\nblog.csdn.net/m0_71369066/article/details/131654824",
             "同类课题选题库;\n对比工作量与创新点定位"],
            ["本科毕业设计答辩经验帖(知乎)\nzhuanlan.zhihu.com/p/521687229",
             "答辩 PPT 与现场答辩实操建议;阶段四答辩准备参考"],
        ], size=9)


# ────────────────────────── 组装 ──────────────────────────
def main():
    doc = Document()
    config_styles(doc)
    build_cover(doc)
    build_front_matter(doc)
    build_body(doc)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("OK saved:", OUT)


if __name__ == "__main__":
    main()
